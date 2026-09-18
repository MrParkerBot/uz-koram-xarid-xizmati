"""Tests for the Excel and PDF export of the list pages (TASK-UZK-042)."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.support import (
    SignedInAdminTestCase,
    a_category,
    a_department,
    a_purchase_application,
    a_supplier,
    an_application,
    an_assigned_application,
    make_user,
    page,
)
from xarid.exports import EXCEL_CONTENT_TYPE, PDF_CONTENT_TYPE
from xarid.models import KATTA_MUTAXASIS, USERS, Contract


def workbook_rows(response) -> list[list[object]]:
    sheet = load_workbook(BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def pdf_text(response) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)


class ExportRouteTests(SignedInAdminTestCase):
    PAGES = ("kelib-arizalar", "qabul-arizalar", "tayinlangan", "kelishinlingan", "xarid-ariza")

    def test_every_list_page_exports_both_formats_with_the_right_headers(self) -> None:
        for page_name in self.PAGES:
            with self.subTest(page=page_name):
                excel = self.client.get(page(f"{page_name}-eksport", "xlsx"))
                pdf = self.client.get(page(f"{page_name}-eksport", "pdf"))

                self.assertEqual(excel.status_code, 200)
                self.assertEqual(excel["Content-Type"], EXCEL_CONTENT_TYPE)
                self.assertIn(f'filename="{page_name}-', excel["Content-Disposition"])
                self.assertEqual(pdf.status_code, 200)
                self.assertEqual(pdf["Content-Type"], PDF_CONTENT_TYPE)
                self.assertIn(".pdf", pdf["Content-Disposition"])

    def test_an_empty_list_exports_headers_and_no_rows(self) -> None:
        excel = self.client.get(page("kelib-arizalar-eksport", "xlsx"))
        pdf = self.client.get(page("kelib-arizalar-eksport", "pdf"))

        rows = workbook_rows(excel)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][:3], ["Ariza raqami", "Bo'lim", "Mahsulot turi"])
        self.assertIn("Ariza raqami", pdf_text(pdf))

    def test_an_unknown_format_is_a_not_found(self) -> None:
        self.assertEqual(self.client.get(page("kelib-arizalar-eksport", "csv")).status_code, 404)

    def test_the_export_answers_to_the_page_permission(self) -> None:
        self.client.logout()
        anonymous = self.client.get(page("kelib-arizalar-eksport", "xlsx"))
        self.assertEqual(anonymous.status_code, 302)
        self.assertTrue(anonymous.url.startswith(reverse("login")))

        self.client.force_login(make_user("requester", user_type=USERS))
        self.assertEqual(self.client.get(page("kelib-arizalar-eksport", "xlsx")).status_code, 403)
        self.assertEqual(self.client.get(page("xarid-ariza-eksport", "xlsx")).status_code, 200)

    def test_the_toolbar_links_carry_the_active_filters(self) -> None:
        department = a_department("Texnik bo`lim")
        an_application(department=department)

        plain = self.client.get(page("kelib-arizalar"))
        filtered = self.client.get(page("kelib-arizalar"), {"bolim": department.pk})

        self.assertContains(plain, f'href="{page("kelib-arizalar-eksport", "xlsx")}"')
        self.assertContains(
            filtered, f'href="{page("kelib-arizalar-eksport", "pdf")}?bolim={department.pk}"'
        )


class ExportContentTests(SignedInAdminTestCase):
    """The rows in the file are the rows on the page, filtered the same way."""

    @classmethod
    def setUpTestData(cls) -> None:
        super().setUpTestData()
        cls.texnik = a_department("Texnik bo`lim")
        cls.moliya = a_department("Moliya bo`limi")
        cls.metal = a_category(100042, "Metallurgiya")
        cls.kimyo = a_category(200031, "Kimyoviy")
        cls.texnik_application = an_application(department=cls.texnik, category=cls.metal)
        cls.moliya_application = an_application(department=cls.moliya, category=cls.kimyo)
        # A second order line makes the export two rows for one application,
        # the way the table renders it.
        cls.moliya_application.items.create(
            mahsulot_turi=cls.metal,
            buyurtma_nomi="Prokat",
            buyurtma_soni="12.5",
            olchov_birligi="kg",
        )

    def test_exporting_with_no_filter_writes_every_table_row(self) -> None:
        excel = self.client.get(page("kelib-arizalar-eksport", "xlsx"))

        rows = workbook_rows(excel)
        self.assertEqual(len(rows), 1 + 3)
        numbers = {row[0] for row in rows[1:]}
        self.assertEqual(
            numbers,
            {self.texnik_application.ariza_raqami, self.moliya_application.ariza_raqami},
        )
        prokat = next(row for row in rows[1:] if row[3] == "Prokat")
        self.assertEqual(prokat[0], self.moliya_application.ariza_raqami)
        self.assertEqual(prokat[1], "Moliya bo`limi")
        self.assertEqual(prokat[4], Decimal("12.5"))

    def test_exporting_with_a_filter_writes_only_the_filtered_rows(self) -> None:
        excel = self.client.get(page("kelib-arizalar-eksport", "xlsx"), {"bolim": self.texnik.pk})

        rows = workbook_rows(excel)
        self.assertEqual(len(rows), 1 + 1)
        self.assertEqual(rows[1][0], self.texnik_application.ariza_raqami)

    def test_the_pdf_carries_the_same_rows_as_the_workbook(self) -> None:
        for query in ({}, {"bolim": self.moliya.pk}):
            with self.subTest(query=query):
                excel = self.client.get(page("kelib-arizalar-eksport", "xlsx"), query)
                pdf = self.client.get(page("kelib-arizalar-eksport", "pdf"), query)

                text = pdf_text(pdf)
                for row in workbook_rows(excel)[1:]:
                    self.assertIn(str(row[0]), text)
                    self.assertIn(str(row[3]), text)
                if query:
                    self.assertNotIn(self.texnik_application.ariza_raqami, text)

    def test_an_invalid_filter_value_exports_the_full_list(self) -> None:
        excel = self.client.get(page("kelib-arizalar-eksport", "xlsx"), {"bolim": "999999"})

        self.assertEqual(len(workbook_rows(excel)), 1 + 3)

    def test_the_assigned_export_is_limited_to_a_specialist_s_own_work(self) -> None:
        specialist = make_user("spec", user_type=KATTA_MUTAXASIS)
        other = make_user("spec2", user_type=KATTA_MUTAXASIS)
        mine = an_assigned_application(self.admin, specialist)
        an_assigned_application(self.admin, other)

        self.client.force_login(specialist)
        rows = workbook_rows(self.client.get(page("tayinlangan-eksport", "xlsx")))

        self.assertEqual(len(rows), 1 + 1)
        self.assertEqual(rows[1][0], mine.ariza_raqami)
        self.assertNotIn("Tayinlangan xodim", rows[0])

        self.client.force_login(self.admin)
        rows = workbook_rows(self.client.get(page("tayinlangan-eksport", "xlsx")))
        self.assertEqual(len(rows), 1 + 2)
        self.assertIn("Tayinlangan xodim", rows[0])

    def test_the_contracts_export_writes_one_row_per_goods_line_with_prices(self) -> None:
        specialist = make_user("spec", user_type=KATTA_MUTAXASIS)
        application = an_assigned_application(self.admin, specialist)
        contract = Contract.raise_contract(
            items=[
                {
                    "buyurtma_nomi": "Bolt",
                    "part_number": "PN-1",
                    "buyurtma_soni": Decimal("2"),
                    "olchov_birligi": "ta",
                    "narxi": Decimal("1000.50"),
                },
                {
                    "buyurtma_nomi": "Gayka",
                    "part_number": "",
                    "buyurtma_soni": Decimal("3"),
                    "olchov_birligi": "ta",
                    "narxi": Decimal("1"),
                },
            ],
            created_by=self.admin,
            application=application,
            supplier=a_supplier(),
        )

        rows = workbook_rows(self.client.get(page("kelishinlingan-eksport", "xlsx")))

        self.assertEqual(len(rows), 1 + 2)
        bolt = next(row for row in rows[1:] if row[2] == "Bolt")
        self.assertEqual(bolt[0], contract.shartnoma_raqami)
        self.assertEqual(bolt[7], Decimal("2001.00"))
        self.assertEqual(bolt[9], "Texnoprom LLC")
        self.assertIn(
            contract.shartnoma_raqami,
            pdf_text(self.client.get(page("kelishinlingan-eksport", "pdf"))),
        )

    def test_the_purchase_export_follows_its_category_filter(self) -> None:
        requester = make_user("requester", user_type=USERS, department=self.texnik)
        kimyo = a_purchase_application(requester, self.texnik, category=self.kimyo, with_pdf=False)
        a_purchase_application(requester, self.texnik, category=self.metal, with_pdf=False)

        rows = workbook_rows(
            self.client.get(page("xarid-ariza-eksport", "xlsx"), {"mahsulot": self.kimyo.pk})
        )

        self.assertEqual(len(rows), 1 + 1)
        self.assertEqual(rows[1][0], kimyo.xarid_raqami)
