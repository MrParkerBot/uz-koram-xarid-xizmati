"""Tests for the Korhona xaridi | Mahsulotlar report (TASK-UZK-047).

One row per ordered product line, with the columns REQ-XARID-003 names, the
filters and period of the other list pages, and a status column that can only
show what the workflow can currently reach.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from tests.support import (
    SignedInAdminTestCase,
    a_category,
    a_contract,
    a_department,
    an_accepted_application,
    an_application,
    arrived_on,
    make_user,
    page,
)
from xarid.models import MENEJER, Application, ApplicationItem, ShartnomaStatus

REFUSED_INVERTED_PERIOD = "Davr boshlanishi tugashidan keyin"


def exported_rows(response) -> list[list[object]]:
    sheet = load_workbook(BytesIO(response.content)).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


class ProductsPageTests(SignedInAdminTestCase):
    """The page itself."""

    @classmethod
    def setUpTestData(cls) -> None:
        super().setUpTestData()
        cls.texnik = a_department("Texnik bo`lim")
        cls.moliya = a_department("Moliya bo`limi")
        cls.metal = a_category(100042, "Metallurgiya")
        cls.kimyo = a_category(200031, "Kimyoviy")

    def test_the_page_answers_with_no_data_at_all(self) -> None:
        response = self.client.get(page("mahsulotlar"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mahsulotlar")
        self.assertContains(response, "Mahsulot topilmadi.")

    def test_an_application_with_two_lines_renders_two_rows(self) -> None:
        application = an_application(department=self.texnik, category=self.metal)
        ApplicationItem.objects.create(
            application=application,
            mahsulot_turi=self.kimyo,
            buyurtma_nomi="Kislota",
            buyurtma_soni=5,
            olchov_birligi="l",
        )

        response = self.client.get(page("mahsulotlar"))

        self.assertEqual(response.content.decode().count('data-id="'), 2)

    def test_every_column_the_requirement_names_is_present(self) -> None:
        an_application(department=self.texnik, category=self.metal)

        response = self.client.get(page("mahsulotlar"))

        # These are literal template text, not values, so they are not escaped.
        for header in (
            "Ariza №",
            "<th>Bo'lim</th>",
            "Mahsulot Turi",
            "Buyurtma nomi",
            "<th>Soni</th>",
            "O'lchov",
            "<th>Izoh</th>",
            "<th>PDF</th>",
            "Qabul sanasi",
            "<th>Holati</th>",
        ):
            with self.subTest(header=header):
                self.assertContains(response, header)

    def test_filtering_by_department_narrows_the_rows(self) -> None:
        texnik = an_application(department=self.texnik, category=self.metal)
        moliya = an_application(department=self.moliya, category=self.metal)

        response = self.client.get(page("mahsulotlar"), {"bolim": self.texnik.pk})

        self.assertContains(response, texnik.ariza_raqami)
        self.assertNotContains(response, moliya.ariza_raqami)

    def test_arriving_from_the_departments_report_narrows_the_page(self) -> None:
        """The Bo`limlar report links here with ?bolim=<id> (TASK-UZK-045)."""
        texnik = an_application(department=self.texnik, category=self.metal)
        moliya = an_application(department=self.moliya, category=self.metal)

        response = self.client.get(f"{page('mahsulotlar')}?bolim={self.texnik.pk}")

        self.assertContains(response, texnik.ariza_raqami)
        self.assertNotContains(response, moliya.ariza_raqami)

    def test_filtering_by_product_type_narrows_the_rows(self) -> None:
        metal = an_application(department=self.texnik, category=self.metal)
        kimyo = an_application(department=self.texnik, category=self.kimyo)

        response = self.client.get(page("mahsulotlar"), {"mahsulot": self.kimyo.pk})

        self.assertContains(response, kimyo.ariza_raqami)
        self.assertNotContains(response, metal.ariza_raqami)

    def test_the_period_narrows_by_the_arrival_date(self) -> None:
        inside = arrived_on(date(2026, 2, 10))
        outside = arrived_on(date(2026, 5, 10))

        response = self.client.get(
            page("mahsulotlar"), {"dan": "2026-02-01", "gacha": "2026-02-28"}
        )

        self.assertContains(response, inside.ariza_raqami)
        self.assertNotContains(response, outside.ariza_raqami)

    def test_an_inverted_period_is_refused_with_a_message(self) -> None:
        application = arrived_on(date(2026, 2, 10))

        response = self.client.get(
            page("mahsulotlar"), {"dan": "2026-05-01", "gacha": "2026-01-01"}
        )

        self.assertContains(response, REFUSED_INVERTED_PERIOD)
        self.assertContains(response, application.ariza_raqami)

    def test_a_type_the_matrix_refuses_cannot_open_it(self) -> None:
        manager = make_user("products.manager", user_type=MENEJER)
        self.client.force_login(manager)

        response = self.client.get(page("mahsulotlar"))

        self.assertEqual(response.status_code, 403)


class ProductsPdfCellTests(SignedInAdminTestCase):
    """The PDF cell follows DEC-019: a link only where one would work."""

    def test_a_row_links_to_the_application_pdf(self) -> None:
        application = an_application(with_pdf=True)

        response = self.client.get(page("mahsulotlar"))

        self.assertContains(response, page("ariza-pdf", application.pk))

    def test_a_row_without_a_pdf_offers_no_link(self) -> None:
        an_application(with_pdf=False)

        response = self.client.get(page("mahsulotlar"))

        self.assertNotContains(response, "ariza-pdf")

    def test_a_rejected_application_offers_no_link(self) -> None:
        """Its attachment is unreachable by DEC-019, so the row must not promise one."""
        application = an_application(with_pdf=True)
        application.reject(by=self.admin, comment="Kerak emas")

        response = self.client.get(page("mahsulotlar"))

        self.assertContains(response, application.ariza_raqami)
        self.assertNotContains(response, page("ariza-pdf", application.pk))


class ProductsStatusTests(SignedInAdminTestCase):
    """What the Holati column can say today."""

    def test_an_application_with_no_contract_shows_its_own_stage(self) -> None:
        application = an_application()

        self.assertEqual(application.current_status_label, "Kelib tushgan")

    def test_an_accepted_application_shows_the_stage_it_reached(self) -> None:
        application = an_accepted_application(self.admin)

        self.assertEqual(application.current_status_label, "Qabul qilingan")

    def test_an_application_with_a_contract_shows_the_contract_status(self) -> None:
        status = ShartnomaStatus.objects.active().first()
        application = an_accepted_application(self.admin)
        a_contract(application, self.admin, status=status)

        self.assertEqual(application.current_status_label, status.name)

    def test_the_most_recent_contract_wins(self) -> None:
        first, second = tuple(ShartnomaStatus.objects.active())[:2]
        application = an_accepted_application(self.admin)
        a_contract(application, self.admin, status=first)
        a_contract(application, self.admin, status=second)

        self.assertEqual(application.current_status_label, second.name)

    def test_a_contract_without_a_status_falls_back_to_the_stage(self) -> None:
        application = an_accepted_application(self.admin)
        a_contract(application, self.admin, status=None)

        self.assertEqual(application.current_status_label, "Qabul qilingan")

    def test_the_page_prints_the_status(self) -> None:
        status = ShartnomaStatus.objects.active().first()
        application = an_accepted_application(self.admin)
        a_contract(application, self.admin, status=status)

        response = self.client.get(page("mahsulotlar"))

        self.assertContains(response, status.name)


class ProductsExportTests(SignedInAdminTestCase):
    """The download holds what the page holds."""

    @classmethod
    def setUpTestData(cls) -> None:
        super().setUpTestData()
        cls.texnik = a_department("Texnik bo`lim")
        cls.moliya = a_department("Moliya bo`limi")

    def test_the_download_has_the_columns_and_one_row_per_line(self) -> None:
        an_application(department=self.texnik)

        rows = exported_rows(self.client.get(page("mahsulotlar-eksport", "xlsx")))

        self.assertEqual(rows[0][0], "Ariza raqami")
        self.assertEqual(rows[0][-1], "Holati")
        self.assertEqual(len(rows), 2)

    def test_the_download_is_narrowed_as_the_page_is(self) -> None:
        texnik = an_application(department=self.texnik)
        an_application(department=self.moliya)

        rows = exported_rows(
            self.client.get(page("mahsulotlar-eksport", "xlsx"), {"bolim": self.texnik.pk})
        )

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], texnik.ariza_raqami)

    def test_both_formats_answer(self) -> None:
        an_application(department=self.texnik)

        for file_format in ("xlsx", "pdf"):
            with self.subTest(file_format=file_format):
                response = self.client.get(page("mahsulotlar-eksport", file_format))

                self.assertEqual(response.status_code, 200)

    def test_an_empty_report_still_downloads_its_headers(self) -> None:
        Application.objects.all().delete()

        rows = exported_rows(self.client.get(page("mahsulotlar-eksport", "xlsx")))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], "Ariza raqami")

    def test_a_type_the_matrix_refuses_cannot_download_either(self) -> None:
        manager = make_user("products.export.manager", user_type=MENEJER)
        self.client.force_login(manager)

        response = self.client.get(page("mahsulotlar-eksport", "xlsx"))

        self.assertEqual(response.status_code, 403)
